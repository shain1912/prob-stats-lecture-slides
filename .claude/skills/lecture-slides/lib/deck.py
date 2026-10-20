"""양식(references/1.pptx) 기반 강의 슬라이드 조립 도구

    from deck import Deck
    d = Deck("references/1.pptx", "lectures/01-introduction/assets")
    d.cover("01", "확률과 통계", "확률변수와 확률분포", "강사 : ○○○", "cover.png")
    s = d.page("1-1", "확률변수란?", "소제목")
    d.save("lectures/01-introduction/output/01-introduction.pptx")

양식 슬라이드 번호: 1 표지, 2 목차, 3 학습 목표, 5 본문(빨간 번호 띠 + ★소제목).
도형 id는 양식 파일 기준 (표지 110/111/114, 목차 123/125/126, 학습 목표 134/135/136, 본문 153/154/156).
"""
import copy
import json
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

INK = RGBColor(0x3F, 0x3F, 0x3F)
RED = RGBColor(0xE2, 0x2B, 0x2D)
GRAY = RGBColor(0x8C, 0x8C, 0x8C)
FONT = "맑은 고딕"


def shape_by_id(shapes, sid):
    for sh in shapes:
        if sh.shape_id == sid:
            return sh
        if sh.shape_type == 6:
            found = shape_by_id(sh.shapes, sid)
            if found is not None:
                return found
    return None


def set_para(para, text):
    """문단의 첫 run 서식을 유지하고 글자만 교체"""
    runs = para.runs
    if not runs:
        para.add_run().text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_paras(shape, texts):
    paras = shape.text_frame.paragraphs
    for i, para in enumerate(paras):
        if i < len(texts):
            set_para(para, texts[i])
        else:
            para._p.getparent().remove(para._p)


def remove(shape):
    shape._element.getparent().remove(shape._element)


def license_label(c):
    return "CC0" if c["license"] == "cc0" else (
        "Public Domain" if c["license"] == "pdm" else f"CC {c['license'].upper()} {c['license_version']}")


class Deck:
    def __init__(self, template, assets):
        self.prs = Presentation(template)
        self.assets = Path(assets)
        self.anim = self.assets / "anim"
        self.photos = self.assets / "photos"
        cred = self.photos / "credits.json"
        self.credits = json.loads(cred.read_text(encoding="utf-8")) if cred.exists() else {}
        self._src = list(self.prs.slides)
        self._cover, self._contents, self._goals, self._page = self._src[0], self._src[1], self._src[2], self._src[4]
        self._moves = []

    # ── 양식 슬라이드 채우기 ────────────────────────────
    def cover(self, number, course, title, lecturer, image=None):
        layout = self._cover.slide_layout
        for sid in (16, 17, 25):                  # 레이아웃의 AI 일러스트 제거
            sh = shape_by_id(layout.shapes, sid)
            if sh is not None:
                remove(sh)
        set_paras(shape_by_id(layout.shapes, 21), [course])
        set_paras(shape_by_id(self._cover.shapes, 110), [number])
        set_paras(shape_by_id(self._cover.shapes, 111), [lecturer])
        t = shape_by_id(self._cover.shapes, 114)
        set_paras(t, [title])
        t.width = Inches(6.3)
        t.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
        if image:
            self.pic(self._cover, self.assets / image, 7.6, 0.9, height=5.9)

    def contents(self, header, items):
        n = len(items)
        set_paras(shape_by_id(self._contents.shapes, 123), sum([[f"{i+1:02d}", ""] for i in range(n)], [])[:-1])
        set_paras(shape_by_id(self._contents.shapes, 125), sum([[x, ""] for x in items], [])[:-1])
        hdr = shape_by_id(self._contents.shapes, 126)
        set_paras(hdr, [header])
        hdr.text_frame.word_wrap = False

    def goals(self, image, goal1, goal2):
        pic = shape_by_id(self._goals.shapes, 134)
        l, t, w, h = pic.left, pic.top, pic.width, pic.height
        remove(pic)
        self._goals.shapes.add_picture(str(self.assets / image), l, t, width=w, height=h)
        set_para(shape_by_id(self._goals.shapes, 135).text_frame.paragraphs[1], goal1)
        set_para(shape_by_id(self._goals.shapes, 136).text_frame.paragraphs[1], goal2)

    def page(self, num, title, subtitle):
        """양식 본문 슬라이드(제목 띠 + ★소제목)를 복제한 새 슬라이드"""
        s = self.prs.slides.add_slide(self._page.slide_layout)
        for ph in list(s.placeholders):
            remove(ph)
        for el in self._page.shapes._spTree.iterchildren():
            tag = el.tag.split("}")[1]
            if tag in ("sp", "grpSp"):
                if tag == "sp" and el.xpath(".//p:cNvPr/@id")[0] == "160":  # 예시 본문 상자 제외
                    continue
                s.shapes._spTree.append(copy.deepcopy(el))
        set_paras(shape_by_id(s.shapes, 153), [num])
        set_paras(shape_by_id(s.shapes, 154), [title])
        set_paras(shape_by_id(s.shapes, 156), [subtitle])
        return s

    # ── 내용 요소 ───────────────────────────────────────
    def text(self, slide, left, top, width, height, items, size=20):
        """items: str(본문 bullet) / ("h", str)(빨간 굵은 소제목) / ("m", str)(들여쓴 수식 줄)"""
        tf = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            kind, text = item if isinstance(item, tuple) else ("b", item)
            run = p.add_run()
            run.text = {"b": "· ", "h": "", "m": "     "}[kind] + text
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.color.rgb = RED if kind == "h" else INK
            run.font.bold = kind == "h"
            p.space_after = Pt(10)
        return tf

    def pic(self, slide, path, left, top, width=None, height=None):
        kw = {}
        if width: kw["width"] = Inches(width)
        if height: kw["height"] = Inches(height)
        return slide.shapes.add_picture(str(path), Inches(left), Inches(top), **kw)

    def photo(self, slide, name, left, top, width=None, height=None):
        """assets/photos/<name>.jpg + 사진 아래 출처 표기"""
        pic = self.pic(slide, self.photos / f"{name}.jpg", left, top, width, height)
        c = self.credits[name]
        tf = slide.shapes.add_textbox(pic.left, pic.top + pic.height, pic.width, Inches(0.3)).text_frame
        run = tf.paragraphs[0].add_run()
        run.text = f"사진: {c['creator']} / {license_label(c)}"
        run.font.name = FONT
        run.font.size = Pt(10)
        run.font.color.rgb = GRAY
        return pic

    def video(self, slide, name, left=1.07, top=1.65, width=11.2):
        """assets/anim/<name>.mp4 (2:1) 삽입, <name>_poster.png를 포스터로 사용.
        기본값은 본문 영역 전체. 위에 정의 문장 2줄을 둘 때는 (2.67, 2.75, 8.6)"""
        w = Inches(width)
        return slide.shapes.add_movie(str(self.anim / f"{name}.mp4"), Inches(left), Inches(top), w, int(w / 2),
                                      poster_frame_image=str(self.anim / f"{name}_poster.png"), mime_type="video/mp4")

    # ── 자주 쓰는 슬라이드 유형 ─────────────────────────
    def analogy(self, num, title, subtitle, photos, items):
        """[비유] 슬라이드: 사진(왼쪽) + 비유 설명(오른쪽). photos: [(name, width_inch)]"""
        s = self.page(num, title, "[비유] " + subtitle)
        x = 1.0
        for name, w in photos:
            self.photo(s, name, x, 1.85, width=w)
            x += w + 0.25
        self.text(s, x + 0.2, 2.0, 12.6 - x, 5.0, items, size=20)
        return s

    def video_page(self, num, title, subtitle, name, lead=None):
        """애니메이션 슬라이드. lead: 영상 위 정의 문장(최대 2줄)"""
        s = self.page(num, title, subtitle)
        if lead:
            self.text(s, 1.4, 1.65, 10.8, 1.0, lead, size=18)
            self.video(s, name, 2.67, 2.75, 8.6)
        else:
            self.video(s, name)
        return s

    def intro_video(self, path, poster, position=1):
        """제목 띠 없는 전체 화면(16:9) 영상 슬라이드. position: 저장 후 슬라이드 위치(0 = 맨 앞, 1 = 표지 다음)"""
        s = self.prs.slides.add_slide(self.prs.slide_layouts[1])          # 양식의 BLANK 레이아웃
        for ph in list(s.placeholders):
            remove(ph)
        s.shapes.add_movie(str(path), 0, 0, self.prs.slide_width, self.prs.slide_height,
                           poster_frame_image=str(poster), mime_type="video/mp4")
        self._moves.append((s, position))
        return s

    def credits_page(self, extra=None):
        """extra: 사진 외 출처 줄 (예: AI 생성 이미지·영상 안내)"""
        if not self.credits and not extra:
            return None
        s = self.page("", "이미지 출처", "사진 (Openverse 검색, CC 라이선스) · AI 생성 이미지·영상")
        lines = [f"{c['title']} — {c['creator']}, {license_label(c)}, {c['source']}" for c in self.credits.values()]
        self.text(s, 1.2, 1.75, 11.2, 5.4, lines + list(extra or []), size=12)
        return s

    def save(self, out):
        """표지·목차·학습 목표 + 새로 만든 슬라이드만 남기고 나머지 양식 슬라이드 삭제"""
        lst = self.prs.slides._sldIdLst
        ids = list(lst)
        keep = ids[:3] + ids[len(self._src):]
        for sid in ids:
            if sid not in keep:
                self.prs.part.drop_rel(sid.rId)
                lst.remove(sid)
        for slide, pos in self._moves:
            sid = next(x for x in lst if self.prs.part.related_part(x.rId) is slide.part)
            lst.remove(sid)
            lst.insert(pos, sid)
        Path(out).parent.mkdir(parents=True, exist_ok=True)
        try:
            self.prs.save(out)
        except PermissionError:
            sys.exit(f"저장 실패: {out} 파일이 PowerPoint에서 열려 있음. 닫고 다시 실행.")
        print("saved", out, len(self.prs.slides), "slides")
